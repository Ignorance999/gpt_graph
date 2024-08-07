import os
import io
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook
from gpt_graph.core.component import Component
import gpt_graph.utils.utils as utils
from pathlib import Path


# %%
class TextExtractor(Component):
    step_type = "node_to_node"
    input_schema = {"input_file_path": {"type": Path}}
    cache_schema = {}
    output_schema = {"extracted_text": {"type": str}}
    output_format = ["plain", "node_like"]

    # def __init__(self):

    def _truncate_text(self, text, word_limit):
        if word_limit is None:
            return text
        words = utils.truncate_text(text, word_limit)
        return words

    def _scrape_pdf(self, word_limit):
        text = ""
        reader = PdfReader(self.input_file_path)
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += " " + page_text
        return self._truncate_text(text, word_limit).strip()

    def _scrape_docx(self, word_limit):
        doc = Document(self.input_file_path)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        return self._truncate_text(text, word_limit)

    def _scrape_xlsx(self):
        workbook = load_workbook(self.input_file_path)
        sheet = workbook.active
        return "\n".join(
            ",".join(str(cell.value) if cell.value is not None else "" for cell in row)
            for row in sheet
        )

    def _scrape_txt(self, word_limit):
        with open(self.input_file_path, "r", encoding="utf-8") as file:
            text = file.read()
        return self._truncate_text(text, word_limit)

    def run(
        self,
        input_file_path,
        output_file_path=None,
        word_limit=None,
        if_save=True,
        if_return_path=False,
    ):
        self.input_file_path = input_file_path
        # Determine the file extension
        file_extension = os.path.splitext(self.input_file_path)[1].lower()
        text = ""

        # Process the file based on its extension
        if file_extension == ".pdf":
            text = self._scrape_pdf(word_limit)
        elif file_extension in [".doc", ".docx"]:
            text = self._scrape_docx(word_limit)
        elif file_extension == ".xlsx":
            text = self._scrape_xlsx(word_limit)
        else:
            text = self._scrape_txt(word_limit)

        # Save the text if requested
        txt_path = None
        if if_save:
            if output_file_path is None:
                output_file_path = os.path.splitext(self.input_file_path)[0] + ".txt"

            with open(output_file_path, "w", encoding="utf-8") as txt_file:
                txt_file.write(text)
                print(f"written at: {output_file_path}")

        if if_return_path:
            text = output_file_path

        # Return the appropriate response
        if self.output_format == "node_like":
            result = {
                "content": text,
                "extra": {
                    "input_file_path": str(self.input_file_path),
                    "output_file_path": str(output_file_path),
                    "word_limit": word_limit,
                    "file_extension": file_extension,
                },
            }
        else:  # Assume "plain" format or any other format
            result = text
        return result


if __name__ == "__main__":
    test_folder = os.environ.get("TEST_FOLDER")
    folder = os.path.join(test_folder, r"\inputs")
    file_path = os.path.join(folder, "accounting.pdf")
    extractor = TextExtractor(output_format="node_like")
    # Example usage:
    extracted_text = extractor.run(
        input_file_path=file_path, word_limit=500, if_save=True, if_return_path=False
    )
    print(extracted_text)
    # # Or, to get the path instead:
    # extracted_file_path = extractor.run(
    #     word_limit=500, if_save=True, if_return_path=True
    # )
    # print(extracted_file_pat
